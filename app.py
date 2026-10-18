import os
import sys
import uuid
import asyncio

if sys.platform == "win32":
    try:
        if sys.stdout:
            sys.stdout.reconfigure(encoding="utf-8")
        if sys.stderr:
            sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass
    asyncio.set_event_loop_policy(asyncio.WindowsProactorEventLoopPolicy())

import io
from typing import Dict, Any, List, Optional
from datetime import datetime

from fastapi import FastAPI, HTTPException, BackgroundTasks
from fastapi.responses import HTMLResponse, StreamingResponse, JSONResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from scraper.maps_scraper import GoogleMapsScraper
from taxonomy import INDUSTRY_TAXONOMY, POPULAR_LOCATIONS

app = FastAPI(title="Google Maps Lead Finder Pro Enterprise", version="2.0.0")

# Mount static files
static_dir = os.path.join(os.path.dirname(__file__), "static")
if not os.path.exists(static_dir):
    os.makedirs(static_dir, exist_ok=True)
app.mount("/static", StaticFiles(directory=static_dir), name="static")

# In-memory store for jobs
class ScrapeJob:
    def __init__(self, job_id: str, business_type: str, location: str, max_results: int, headless: bool, search_mode: str = "standard"):
        self.job_id = job_id
        self.business_type = business_type
        self.location = location
        self.max_results = max_results
        self.headless = headless
        self.search_mode = search_mode
        self.status = "running"  # running, completed, stopped, error
        self.leads: List[Dict[str, Any]] = []
        self.logs: List[str] = []
        self.queue: asyncio.Queue = asyncio.Queue()
        self.scraper: Optional[GoogleMapsScraper] = None
        self.created_at = datetime.now()

jobs: Dict[str, ScrapeJob] = {}

class ScrapeRequest(BaseModel):
    business_type: str = Field(..., min_length=1, description="e.g. Textile Manufacturer, Kirana Store, Dentist")
    location: str = Field(..., min_length=1, description="e.g. Kanpur, Jaipur, New York, London, Dubai")
    max_results: int = Field(20, ge=1, le=100)
    headless: bool = Field(True, description="True for background, False to view browser")
    search_mode: str = Field("standard", description="standard | industrial | small_stores")

@app.get("/", response_class=HTMLResponse)
async def serve_index():
    index_file = os.path.join(static_dir, "index.html")
    if os.path.exists(index_file):
        with open(index_file, "r", encoding="utf-8") as f:
            return HTMLResponse(content=f.read())
    return HTMLResponse("<h1>Lead Finder Pro Enterprise Loading... Please refresh.</h1>")

@app.get("/sw.js")
async def serve_sw():
    sw_file = os.path.join(static_dir, "sw.js")
    if os.path.exists(sw_file):
        return FileResponse(sw_file, media_type="application/javascript")
    raise HTTPException(status_code=404, detail="Service worker not found")

@app.get("/manifest.json")
async def serve_manifest():
    manifest_file = os.path.join(static_dir, "manifest.json")
    if os.path.exists(manifest_file):
        return FileResponse(manifest_file, media_type="application/manifest+json")
    raise HTTPException(status_code=404, detail="Manifest not found")

@app.get("/api/taxonomy")
async def get_taxonomy():
    """Returns curated industries, manufacturing sectors, and global locations."""
    return {
        "industries": INDUSTRY_TAXONOMY,
        "locations": POPULAR_LOCATIONS
    }

async def run_scraper_task(job: ScrapeJob):
    scraper = GoogleMapsScraper(headless=job.headless)
    job.scraper = scraper
    
    try:
        async for event in scraper.scrape(
            business_type=job.business_type,
            location=job.location,
            max_results=job.max_results,
            search_mode=job.search_mode
        ):
            event_type = event.get("type")
            if event_type == "status":
                msg = event["message"]
                job.logs.append(msg)
                await job.queue.put({"type": "log", "message": msg})
            elif event_type == "lead":
                lead = event["data"]
                job.leads.append(lead)
                await job.queue.put({"type": "lead", "data": lead})
            elif event_type == "complete":
                job.status = "completed"
                stats = event.get("stats", {})
                await job.queue.put({"type": "complete", "stats": stats, "total": len(job.leads)})
                break
            elif event_type == "error":
                job.status = "error"
                err_msg = event.get("message", "Unknown error")
                job.logs.append(f"ERROR: {err_msg}")
                await job.queue.put({"type": "error", "message": err_msg})
                break
    except Exception as e:
        job.status = "error"
        err_msg = str(e)
        job.logs.append(f"FATAL: {err_msg}")
        await job.queue.put({"type": "error", "message": err_msg})
    finally:
        await job.queue.put({"type": "done"})

@app.post("/api/scrape/start")
async def start_scraping(req: ScrapeRequest, background_tasks: BackgroundTasks):
    job_id = str(uuid.uuid4())[:8]
    job = ScrapeJob(
        job_id=job_id,
        business_type=req.business_type.strip(),
        location=req.location.strip(),
        max_results=req.max_results,
        headless=req.headless,
        search_mode=req.search_mode
    )
    jobs[job_id] = job
    background_tasks.add_task(run_scraper_task, job)
    return {
        "status": "success",
        "job_id": job_id,
        "message": f"Scraping started for '{job.business_type}' in '{job.location}' [{job.search_mode}]"
    }

@app.get("/api/scrape/stream/{job_id}")
async def stream_scrape_events(job_id: str):
    if job_id not in jobs:
        raise HTTPException(status_code=404, detail="Job not found")
    
    job = jobs[job_id]

    async def event_generator():
        import json
        # Send existing leads if reconnecting
        for lead in job.leads:
            yield f"event: lead\ndata: {json.dumps(lead)}\n\n"

        while True:
            try:
                event = await asyncio.wait_for(job.queue.get(), timeout=20.0)
                if event.get("type") == "done":
                    yield f"event: done\ndata: {json.dumps({'status': job.status, 'total': len(job.leads)})}\n\n"
                    break

                event_name = event.get("type", "message")
                if event_name == "lead":
                    lead_data = event.get("data", event)
                    yield f"event: lead\ndata: {json.dumps(lead_data)}\n\n"
                elif event_name == "log":
                    yield f"event: log\ndata: {json.dumps({'message': event.get('message', '')})}\n\n"
                elif event_name == "complete":
                    yield f"event: complete\ndata: {json.dumps(event)}\n\n"
                else:
                    yield f"event: {event_name}\ndata: {json.dumps(event)}\n\n"
            except asyncio.TimeoutError:
                yield ": keepalive\n\n"
            except asyncio.CancelledError:
                break

    return StreamingResponse(
        event_generator(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no"
        }
    )

@app.post("/api/scrape/stop/{job_id}")
async def stop_scraping(job_id: str):
    if job_id not in jobs:
        raise HTTPException(status_code=404, detail="Job not found")
    job = jobs[job_id]
    if job.scraper:
        job.scraper.stop()
    job.status = "stopped"
    return {"status": "success", "message": "Scraping cancellation initiated"}

@app.get("/api/job/{job_id}")
async def get_job_status(job_id: str):
    if job_id not in jobs:
        raise HTTPException(status_code=404, detail="Job not found")
    job = jobs[job_id]
    no_web_count = sum(1 for l in job.leads if l.get("has_website") == "NO")
    has_web_count = len(job.leads) - no_web_count
    unclaimed_count = sum(1 for l in job.leads if "Unclaimed" in str(l.get("profile_status", "")))
    with_email_count = sum(1 for l in job.leads if "@" in str(l.get("email", "")))

    return {
        "job_id": job.job_id,
        "business_type": job.business_type,
        "location": job.location,
        "status": job.status,
        "total_leads": len(job.leads),
        "missing_website_count": no_web_count,
        "has_website_count": has_web_count,
        "unclaimed_count": unclaimed_count,
        "with_email_count": with_email_count,
        "leads": job.leads,
        "logs": job.logs[-25:]
    }

def format_lead_sheet(ws, title, leads, header_fill, header_font, yes_fill, yes_font, no_fill, no_font, link_font, regular_font, thin_border):
    ws.title = title
    headers = [
        "S.No",
        "Business Name",
        "Category",
        "Phone Number",
        "WhatsApp Link",
        "Email Address",
        "Has Website?",
        "Website URL",
        "Profile Status",
        "Rating",
        "Reviews Count",
        "Full Address",
        "Google Maps URL"
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.row_dimensions[1].height = 28

    for i, lead in enumerate(leads, start=1):
        row_num = i + 1
        has_web = lead.get("has_website", "NO")
        web_url = lead.get("website", "")
        if web_url == "None (Opportunity)":
            web_url = ""

        wa_url = lead.get("whatsapp_url", "")
        email = lead.get("email", "")
        if email in ("Scanning...", "Not Available"):
            email = ""

        row_data = [
            i,
            lead.get("name", ""),
            lead.get("category", ""),
            lead.get("phone", ""),
            "Chat on WhatsApp" if wa_url else "-",
            email,
            has_web,
            web_url,
            lead.get("profile_status", "Claimed"),
            lead.get("rating", ""),
            lead.get("reviews_count", ""),
            lead.get("address", ""),
            lead.get("maps_url", "")
        ]
        ws.append(row_data)

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=row_num, column=col_idx)
            c.font = regular_font
            c.border = thin_border
            c.alignment = Alignment(vertical="center")

            # WhatsApp Link
            if col_idx == 5 and wa_url:
                c.hyperlink = wa_url
                c.font = Font(name="Calibri", size=11, color="16A34A", underline="single", bold=True)

            # Email Link
            if col_idx == 6 and email:
                c.hyperlink = f"mailto:{email}"
                c.font = link_font

            # Has Website formatting
            if col_idx == 7:
                if has_web == "YES":
                    c.fill = yes_fill
                    c.font = yes_font
                    c.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    c.fill = no_fill
                    c.font = no_font
                    c.alignment = Alignment(horizontal="center", vertical="center")

            # Website Link formatting
            if col_idx == 8 and web_url and web_url.startswith("http"):
                c.hyperlink = web_url
                c.font = link_font

            # Maps Link formatting
            if col_idx == 13 and lead.get("maps_url", "").startswith("http"):
                c.hyperlink = lead.get("maps_url")
                c.value = "View on Maps"
                c.font = link_font

        ws.row_dimensions[row_num].height = 22

    # Column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 48)

@app.get("/api/export/{job_id}/excel")
@app.get("/api/scrape/export/excel/{job_id}")
async def export_excel(job_id: str):
    if job_id not in jobs:
        raise HTTPException(status_code=404, detail="Job not found")
    job = jobs[job_id]
    if not job.leads:
        raise HTTPException(status_code=400, detail="No leads collected for this job yet.")

    wb = openpyxl.Workbook()

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    yes_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    yes_font = Font(name="Calibri", size=11, bold=True, color="065F46")
    no_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    no_font = Font(name="Calibri", size=11, bold=True, color="991B1B")
    link_font = Font(name="Calibri", size=11, color="2563EB", underline="single")
    regular_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Sheet 1: All Leads
    ws1 = wb.active
    format_lead_sheet(ws1, "All Leads", job.leads, header_fill, header_font, yes_fill, yes_font, no_fill, no_font, link_font, regular_font, thin_border)

    # Sheet 2: Missing Website (Hot Prospects)
    hot_leads = [l for l in job.leads if l.get("has_website") == "NO"]
    if hot_leads:
        ws2 = wb.create_sheet(title="No Website (Hot Leads)")
        format_lead_sheet(ws2, "No Website (Hot Leads)", hot_leads, header_fill, header_font, yes_fill, yes_font, no_fill, no_font, link_font, regular_font, thin_border)

    # Sheet 3: Unclaimed Profiles
    unclaimed = [l for l in job.leads if "Unclaimed" in str(l.get("profile_status", ""))]
    if unclaimed:
        ws3 = wb.create_sheet(title="Unclaimed GMB Profiles")
        format_lead_sheet(ws3, "Unclaimed GMB Profiles", unclaimed, header_fill, header_font, yes_fill, yes_font, no_fill, no_font, link_font, regular_font, thin_border)

    # Sheet 4: With Email
    with_email = [l for l in job.leads if "@" in str(l.get("email", ""))]
    if with_email:
        ws4 = wb.create_sheet(title="Verified Emails")
        format_lead_sheet(ws4, "Verified Emails", with_email, header_fill, header_font, yes_fill, yes_font, no_fill, no_font, link_font, regular_font, thin_border)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    clean_business = "".join(c for c in job.business_type if c.isalnum() or c in (' ', '_')).rstrip().replace(' ', '_')
    clean_loc = "".join(c for c in job.location if c.isalnum() or c in (' ', '_')).rstrip().replace(' ', '_')
    filename = f"leads_{clean_business}_{clean_loc}.xlsx"

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@app.get("/api/export/{job_id}/csv")
@app.get("/api/scrape/export/csv/{job_id}")
async def export_csv(job_id: str):
    if job_id not in jobs:
        raise HTTPException(status_code=404, detail="Job not found")
    job = jobs[job_id]
    if not job.leads:
        raise HTTPException(status_code=400, detail="No leads collected for this job yet.")

    df = pd.DataFrame(job.leads)
    rename_map = {
        "index": "S.No",
        "name": "Business Name",
        "category": "Category",
        "phone": "Phone Number",
        "whatsapp_url": "WhatsApp Direct Link",
        "email": "Email Address",
        "has_website": "Has Website?",
        "website": "Website URL",
        "profile_status": "Google Profile Status",
        "rating": "Rating",
        "reviews_count": "Reviews Count",
        "address": "Full Address",
        "maps_url": "Google Maps Link"
    }
    df = df.rename(columns={k: v for k, v in rename_map.items() if k in df.columns})

    csv_buffer = io.StringIO()
    df.to_csv(csv_buffer, index=False, encoding="utf-8-sig")
    csv_buffer.seek(0)

    clean_business = "".join(c for c in job.business_type if c.isalnum() or c in (' ', '_')).rstrip().replace(' ', '_')
    clean_loc = "".join(c for c in job.location if c.isalnum() or c in (' ', '_')).rstrip().replace(' ', '_')
    filename = f"leads_{clean_business}_{clean_loc}.csv"

    return StreamingResponse(
        iter([csv_buffer.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

if __name__ == "__main__":
    import uvicorn
    print("Starting Google Maps Lead Finder Pro Enterprise on http://127.0.0.1:8000 ...")
    uvicorn.run("app:app", host="0.0.0.0", port=8000, reload=True)
